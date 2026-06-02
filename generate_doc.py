from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# Page margins
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1.2)
section.right_margin = Inches(1.2)

def add_heading(doc, text, size=16, bold=True, color=RGBColor(0x1F, 0x49, 0x7D), center=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def add_body(doc, text, size=11, bold=False, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Inches(0.3)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    return p

def add_divider(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F497D')
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(4)

# ─── TITLE ───
add_heading(doc, "PROMET Project", size=20, center=True)
add_heading(doc, "Automated Android Application Security Testing Platform", size=13, bold=False,
            color=RGBColor(0x44, 0x72, 0xC4), center=True)
doc.add_paragraph()
add_divider(doc)
doc.add_paragraph()

# ─── METHOD SUMMARY ───
add_heading(doc, "Method Summary", size=13)
add_body(doc,
    "The concept is based on integrating an AI agent with a set of specialized security tools, "
    "where the agent takes full control of the Android application analysis process — "
    "from decompilation through vulnerability discovery and documentation — in an automated and structured manner.",
    size=11)

doc.add_paragraph()

# ─── PROPOSED STEPS ───
add_heading(doc, "Proposed Steps", size=13)

steps = [
    ("1.", "Set up a rooted Android emulator environment and automate its initialization."),
    ("2.", "Static analysis of the application (decompilation, component mapping, automated scanning)."),
    ("3.", "Dynamic analysis (traffic interception, memory manipulation at runtime)."),
    ("4.", "Log findings into a structured database with supporting evidence."),
    ("5.", "Generate a final report listing discovered vulnerabilities and remediation recommendations."),
]

for num, text in steps:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.3)
    run_num = p.add_run(num + " ")
    run_num.bold = True
    run_num.font.size = Pt(11)
    run_body = p.add_run(text)
    run_body.font.size = Pt(11)

doc.add_paragraph()
add_divider(doc)
doc.add_paragraph()

# ─── PRACTICAL EXAMPLE ───
add_heading(doc, "Practical Example — Banking Application", size=13)
add_body(doc, "Input:  bank_app.apk", size=11, bold=True)
doc.add_paragraph()

table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'

hdr = table.rows[0].cells
hdr[0].text = "Phase"
hdr[1].text = "What the System Does"
for cell in hdr:
    cell.paragraphs[0].runs[0].bold = True
    cell.paragraphs[0].runs[0].font.size = Pt(11)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), '1F497D')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

rows_data = [
    ("1", "Decompiles the application and inspects its internal components."),
    ("2", "Discovers that the app stores the login Token in an unencrypted location on the device."),
    ("3", "Intercepts server communication and detects that data is sent without certificate validation."),
    ("4", "Proves the vulnerability with executable code and captures the evidence."),
    ("5", "Records the result and generates the final security report."),
]

for phase, action in rows_data:
    row = table.add_row().cells
    row[0].text = phase
    row[1].text = action
    row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for cell in row:
        cell.paragraphs[0].runs[0].font.size = Pt(10)

doc.add_paragraph()
add_heading(doc, "Output:", size=11, bold=True, color=RGBColor(0x1F, 0x49, 0x7D))

output_lines = [
    "Vulnerability: Unencrypted Token Storage          →  Severity: High",
    "Vulnerability: Missing Certificate Validation     →  Severity: Critical",
    "Evidence: Screenshot + intercepted network data",
    "Recommendation: Encrypt local storage + Enable Certificate Pinning",
]
for line in output_lines:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    run = p.add_run(line)
    run.font.size = Pt(10)
    run.font.name = "Courier New"

doc.add_paragraph()
add_divider(doc)
doc.add_paragraph()

# ─── IMPORTANT NOTE ───
add_heading(doc, "Important Note", size=13)
add_body(doc,
    "The proposed steps are subject to change. Given the rapid advancement in the field of Artificial Intelligence, "
    "the most up-to-date methodology will be determined and applied at each stage of implementation "
    "as the project progresses.",
    size=11)

doc.save("/home/user/PROMET/PROMET_Summary.docx")
print("Done.")
